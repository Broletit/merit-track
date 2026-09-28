
from pathlib import Path
import re, math, textwrap, zipfile, importlib.util

import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, Ellipse, Circle, Polygon, FancyArrowPatch

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path("thesis/generated_chapter3")
IMG = ROOT / "diagrams"
XMI = ROOT / "visual-paradigm"
EDIT = ROOT / "editable-diagrams"
DRAWIO_UC = EDIT / "usecase-drawio"
DRAWIO_BP = EDIT / "business-process-drawio"
ACTIVITY_XMI = EDIT / "business-process-xmi"
ROOT.mkdir(parents=True, exist_ok=True)
IMG.mkdir(parents=True, exist_ok=True)
XMI.mkdir(parents=True, exist_ok=True)
DRAWIO_UC.mkdir(parents=True, exist_ok=True)
DRAWIO_BP.mkdir(parents=True, exist_ok=True)
ACTIVITY_XMI.mkdir(parents=True, exist_ok=True)


# Data is embedded so the generator is self-contained.
ACTORS = [
    ("Sinh viên", "Tham gia hoạt động, theo dõi điểm rèn luyện, tham gia đợt xét và quản lý hồ sơ xét sinh viên."),
    ("Cán bộ lớp", "Duyệt hồ sơ sinh viên vòng 1; đồng thời tham gia hoạt động và đợt xét dành cho cán bộ."),
    ("Cán bộ khoa", "Theo dõi các lớp được phân công, điểm danh hoạt động, duyệt hồ sơ vòng 2; đồng thời tham gia nghiệp vụ cán bộ."),
    ("Quản trị viên", "Quản trị học kỳ, người dùng, lớp, hoạt động, khung điểm rèn luyện, bộ tiêu chuẩn, đợt xét, hồ sơ cán bộ, báo cáo và nhật ký."),
]
ROLE_MATRIX = []
BUSINESS_RULES = [
    ("BR01", "Tại một thời điểm chỉ có một học kỳ hiện hành; activity/event mới sử dụng học kỳ hiện hành."),
    ("BR02", "Activity/Event của học kỳ cũ được giới hạn thao tác thay đổi và chủ yếu dùng để tra cứu lịch sử."),
    ("BR03", "Thời gian đăng ký và thời gian diễn ra hoạt động phải nằm trong khoảng thời gian học kỳ."),
    ("BR04", "Người dùng chỉ đăng ký activity đã published, đúng audience, đúng thời gian, đúng scope, không khóa và có participation_source phù hợp."),
    ("BR05", "Hủy đăng ký chỉ thực hiện khi registration còn registered, chưa hết hạn đăng ký, activity chưa bắt đầu và không bị khóa."),
    ("BR06", "Điểm danh thành công yêu cầu activity published, đang diễn ra, user hợp lệ, đã đăng ký và chưa attended."),
    ("BR07", "Điểm rèn luyện từ activity chỉ được ghi nhận khi điểm danh thành công; điểm thực nhận bị giới hạn bởi khung điểm."),
    ("BR08", "Template đã được event sử dụng không được cấu hình trực tiếp; cần clone để tạo phiên bản mới."),
    ("BR09", "event.type phải khớp criteria_template.for_type."),
    ("BR10", "Snapshot tiêu chuẩn của event chỉ được đồng bộ khi event chưa có submission."),
    ("BR11", "Auto criterion theo activity chỉ đạt khi activity_registrations.status='attended'."),
    ("BR12", "Student submission đi qua hai vòng: cán bộ lớp vòng 1, cán bộ khoa vòng 2."),
    ("BR13", "Officer submission do Admin duyệt một vòng."),
    ("BR14", "Faculty officer chỉ duyệt vòng 2 cho lớp có faculty_class_assignments tương ứng."),
    ("BR15", "Class officer chỉ duyệt vòng 1 cho hồ sơ thuộc lớp mà mình còn là thành viên hiện hành."),
    ("BR16", "Tài khoản class_officer/faculty_officer mặc định đăng nhập ở student context và có thể chuyển context theo role."),
    ("BR17", "Minh chứng sinh viên được kiểm tra loại tệp và dung lượng trước khi lưu; source hiện hỗ trợ PDF/PNG tối đa 10 MB."),
    ("BR18", "Khung điểm rèn luyện có cấu trúc cha-con; điểm của từng mục bị giới hạn bởi score_max."),
    ("BR19", "Officer score_total chỉ được tính cho event type=officer; student submission có score_total=0."),
    ("BR20", "Các thay đổi quản trị nhạy cảm có thể yêu cầu lý do và được ghi audit_logs."),
]
DB_TABLES = []

def mk(id,name,actor,secondary,desc,pre,post,main,alt=None,exc=None):
    return dict(id=id,name=name,actor=actor,secondary=secondary,desc=desc,pre=pre,post=post,main=main,alt=alt or ["Không có."],exc=exc or ["Không có."])

def pair(actor_action, system_action):
    return (actor_action, system_action)

def simple_view(name, actor, obj, scope="phù hợp"):
    return [
        pair(f"Chọn chức năng “{name}”.", f"Tiếp nhận yêu cầu và xác định phạm vi dữ liệu {scope}."),
        pair("", f"Truy vấn {obj} theo quyền và bộ lọc hiện hành."),
        pair("Chọn bộ lọc hoặc một bản ghi để xem chi tiết (nếu cần).", f"Hiển thị {obj} tương ứng.")
    ]

USECASES = [
mk("UC01","Đăng nhập","Người dùng","Không","Xác thực tài khoản bằng MSSV/email và mật khẩu để truy cập MeritTrack.","Đã có tài khoản.","Tạo session và mở dashboard phù hợp.",
   [pair("Nhập MSSV/email và mật khẩu, chọn Đăng nhập.","Kiểm tra trường bắt buộc, tìm tài khoản và xác minh mật khẩu."),
    pair("", "Kiểm tra trạng thái tài khoản và xác định login_context mặc định."),
    pair("", "Tạo session và hiển thị dashboard theo context.")],
   ["Tài khoản cán bộ đăng nhập mặc định ở student context và có thể chuyển context sau đó."],
   ["Thông tin đăng nhập không hợp lệ: hiển thị lỗi và cho phép nhập lại.","Tài khoản bị khóa: từ chối đăng nhập."]),
mk("UC02","Đổi mật khẩu","Người dùng","Không","Thay đổi mật khẩu đang sử dụng.","Đã đăng nhập.","Mật khẩu mới được lưu.",
   [pair("Mở chức năng Đổi mật khẩu.","Hiển thị biểu mẫu đổi mật khẩu."),
    pair("Nhập mật khẩu hiện tại, mật khẩu mới và xác nhận; chọn Lưu.","Kiểm tra mật khẩu hiện tại và tính hợp lệ của mật khẩu mới."),
    pair("", "Cập nhật mật khẩu và thông báo thành công.")],
   exc=["Mật khẩu hiện tại sai hoặc xác nhận không khớp: hiển thị lỗi và quay lại bước nhập."]),
mk("UC03","Chuyển ngữ cảnh sử dụng","Cán bộ lớp, Cán bộ khoa","Không","Chuyển giữa giao diện người dùng và giao diện cán bộ trên cùng tài khoản.","Đã đăng nhập và role hỗ trợ context yêu cầu.","login_context được cập nhật.",
   [pair("Chọn ngữ cảnh muốn sử dụng.","Kiểm tra role có quyền sử dụng context đã chọn."),
    pair("", "Cập nhật session.login_context và hiển thị dashboard tương ứng.")],
   ["Class officer: student↔class_officer; Faculty officer: student↔faculty_officer."],
   ["Context không hợp lệ: giữ nguyên context hiện tại và thông báo lỗi."]),
mk("UC04","Đăng xuất","Người dùng","Không","Kết thúc phiên đăng nhập.","Có session.","Session kết thúc.",
   [pair("Chọn Đăng xuất.","Thu hồi session và chuyển về trang đăng nhập.")]),
mk("UC05","Xem thông tin cá nhân","Sinh viên, Cán bộ lớp, Cán bộ khoa","Không","Xem thông tin tài khoản và lớp hiện hành.","Đã đăng nhập.","Thông tin được hiển thị.",
   simple_view("Thông tin cá nhân","Sinh viên","thông tin tài khoản và lớp hiện hành")),
mk("UC06","Hiển thị QR cá nhân","Sinh viên, Cán bộ lớp, Cán bộ khoa","Cán bộ khoa/Admin","Hiển thị QR động phục vụ điểm danh.","Đã đăng nhập; không phải admin.","QR hợp lệ được hiển thị.",
   [pair("Chọn Mở QR điểm danh.","Tạo payload QR từ user id và MSSV, ký payload và hiển thị QR."),
    pair("Giữ hộp QR mở trong lúc điểm danh.","Làm mới payload QR theo chu kỳ của giao diện.")],
   exc=["Không tạo được QR: hiển thị lỗi và cho phép thử lại."]),
mk("UC07","Xem thông báo","Người dùng","Không","Xem thông báo nghiệp vụ.","Đã đăng nhập.","Danh sách thông báo được hiển thị.",
   simple_view("Thông báo","Người dùng","notifications của tài khoản")),
mk("UC08","Đánh dấu thông báo đã đọc","Người dùng","Không","Cập nhật trạng thái đã đọc.","Đã đăng nhập; notification thuộc user.","is_read/read_at được cập nhật.",
   [pair("Chọn đánh dấu đã đọc.","Kiểm tra notification thuộc tài khoản."),
    pair("", "Cập nhật trạng thái đọc và làm mới danh sách.")],
   exc=["Notification không thuộc tài khoản: không cập nhật."]),
mk("UC09","Xem danh sách hoạt động","Sinh viên","Không","Xem activity dành cho sinh viên theo học kỳ/phạm vi.","Student context.","Danh sách activity được hiển thị.",
   simple_view("Hoạt động","Sinh viên","danh sách activity published","theo học kỳ và scope lớp")),
mk("UC10","Xem chi tiết hoạt động","Sinh viên","Không","Xem thông tin activity và trạng thái đăng ký.","Activity thuộc phạm vi.","Chi tiết và quyền đăng ký/hủy được hiển thị.",
   [pair("Chọn một hoạt động.","Kiểm tra activity, audience, status và scope."),
    pair("", "Đọc registration hiện tại và tính canRegister/canCancel."),
    pair("", "Hiển thị chi tiết hoạt động.")],
   exc=["Activity không tồn tại hoặc ngoài phạm vi: không hiển thị chi tiết."]),
mk("UC11","Đăng ký hoạt động","Sinh viên","Không","Đăng ký activity nội bộ.","Activity published, đúng thời gian/audience/scope, không khóa.","Registration=registered.",
   [pair("Chọn Đăng ký.","Kiểm tra status, thời gian, audience, lớp/scope, khóa và participation_source."),
    pair("", "Kiểm tra đăng ký hiện có."),
    pair("", "Tạo mới hoặc kích hoạt lại registration=registered và thông báo thành công.")],
   ["Nếu registration trước đó cancelled và vẫn đủ điều kiện, kích hoạt lại registration."],
   ["Không đủ điều kiện đăng ký: hiển thị lý do và giữ nguyên dữ liệu."]),
mk("UC12","Hủy đăng ký hoạt động","Sinh viên","Không","Hủy registration trước activity.","Registration=registered; còn hạn; chưa bắt đầu; không khóa.","Registration=cancelled.",
   [pair("Chọn Hủy đăng ký.","Kiểm tra registration, hạn đăng ký, start_at và registration_locked."),
    pair("Xác nhận hủy.","Cập nhật status=cancelled và thông báo thành công.")],
   ["Nếu còn thời gian, có thể đăng ký lại."],
   ["Không còn đủ điều kiện hủy: thông báo và giữ nguyên registration."]),
mk("UC13","Xem điểm rèn luyện","Sinh viên","Không","Theo dõi điểm ĐRL theo học kỳ và khung phân cấp.","Student context.","Điểm tổng hợp được hiển thị.",
   [pair("Chọn Điểm rèn luyện.","Xác định học kỳ và đọc khung conduct_score_categories."),
    pair("", "Đọc conduct_scores, tính điểm từ node con lên cha và cap score_max."),
    pair("Chọn học kỳ hoặc bộ lọc khác (nếu cần).","Hiển thị tổng điểm, điểm từng mục và nguồn điểm.")]),
mk("UC14","Xem danh sách đợt xét sinh viên","Sinh viên","Không","Xem events type=student.","Student context.","Danh sách event và trạng thái submission được hiển thị.",
   simple_view("Đợt xét","Sinh viên","events type=student","theo học kỳ và phạm vi")),
mk("UC15","Xem chi tiết đợt xét sinh viên","Sinh viên","Không","Xem event và snapshot criteria.","Event student thuộc phạm vi.","Tiêu chuẩn/tiêu chí và thao tác hồ sơ được hiển thị.",
   [pair("Chọn một đợt xét.","Đọc event, event_criteria_groups/items và submission hiện có."),
    pair("", "Tính phase nhận hồ sơ và khả năng tạo/chỉnh sửa/nộp."),
    pair("", "Hiển thị tiêu chuẩn, tiêu chí và trạng thái hồ sơ.")]),
mk("UC16","Khởi tạo hồ sơ sinh viên","Sinh viên","Không","Tạo hoặc tiếp tục submission draft.","Event student published, đang nhận hồ sơ; user có lớp.","Submission draft tồn tại.",
   [pair("Chọn Bắt đầu/tiếp tục hồ sơ.","Kiểm tra event phase và submission hiện có."),
    pair("", "Nếu chưa có, xác định class_id và tạo submission=draft."),
    pair("", "Chạy auto-evaluation và mở trang hồ sơ.")],
   ["Nếu đã có submission editable, mở lại hồ sơ đó."],
   ["Event không nhận hồ sơ hoặc user chưa có lớp: không tạo submission."]),
mk("UC17","Bổ sung minh chứng sinh viên","Sinh viên","Không","Lưu note/file minh chứng.","Submission thuộc user và editable.","Evidence được lưu.",
   [pair("Chọn tiêu chí cần bổ sung minh chứng.","Hiển thị trường note/tệp theo evidence_type."),
    pair("Nhập ghi chú hoặc chọn tệp; chọn Lưu.","Kiểm tra quyền, phase, định dạng và dung lượng."),
    pair("", "Lưu submission_items/files, cập nhật updated_at và chạy auto-evaluation.")],
   exc=["Tệp sai loại/dung lượng hoặc hồ sơ hết hạn: hiển thị lỗi và không lưu."]),
mk("UC18","Yêu cầu hỗ trợ hồ sơ","Sinh viên","Admin/Khoa","Gửi yêu cầu hỗ trợ submission.","Submission thuộc user và editable.","Support request=open.",
   [pair("Mở chức năng Yêu cầu hỗ trợ.","Hiển thị biểu mẫu nhập nội dung."),
    pair("Nhập nội dung và chọn Gửi.","Kiểm tra độ dài và quyền sở hữu submission."),
    pair("", "Tạo/cập nhật support request=open và thông báo thành công.")],
   exc=["Nội dung không hợp lệ hoặc submission không editable: không tạo yêu cầu."]),
mk("UC19","Nộp hồ sơ sinh viên","Sinh viên","Cán bộ lớp","Gửi hồ sơ vào vòng xét 1.","Submission thuộc user; event đang nhận; status cho phép.","submitted_v1.",
   [pair("Chọn Nộp hồ sơ.","Kiểm tra quyền, status và event phase."),
    pair("", "Chạy auto-evaluation và kiểm tra required/min_required."),
    pair("Xác nhận nộp hồ sơ.","Cập nhật submitted_v1, submitted_at và timeline; đưa hồ sơ vào queue vòng 1.")],
   ["Sau needs_revision_v1, chỉnh sửa rồi nộp lại khi còn thời gian."],
   ["Thiếu điều kiện bắt buộc: hiển thị tiêu chí chưa đạt và quay lại chỉnh sửa."]),
mk("UC20","Theo dõi hồ sơ sinh viên","Sinh viên","Không","Theo dõi status, evidence, review và timeline.","Có submission thuộc user.","Tiến trình được hiển thị.",
   simple_view("Hồ sơ","Sinh viên","submission, evidence, auto result, reviews và timeline")),
mk("UC21","Xem hồ sơ chờ duyệt vòng 1","Cán bộ lớp","Sinh viên","Xem queue submitted_v1 của lớp.","Class officer context và còn thuộc lớp.","Queue được hiển thị.",
   simple_view("Duyệt vòng 1","Cán bộ lớp","submission submitted_v1","thuộc lớp hiện hành của reviewer")),
mk("UC22","Đánh giá tiêu chí vòng 1","Cán bộ lớp","Sinh viên","Đánh giá tiêu chí thủ công ở vòng 1.","submitted_v1; đúng lớp; có evidence; không auto-pass.","criteria review round=1.",
   [pair("Mở một tiêu chí cần đánh giá.","Kiểm tra submission, auto result, evidence và quyền lớp."),
    pair("Chọn Đạt/Chưa đạt và nhập ghi chú (nếu có).","Upsert submission_criteria_reviews round=1."),
    pair("", "Hiển thị trạng thái đánh giá mới.")],
   exc=["Tiêu chí auto-pass, thiếu evidence hoặc sai quyền: từ chối cập nhật."]),
mk("UC23","Duyệt hồ sơ vòng 1","Cán bộ lớp","Sinh viên","Chốt quyết định vòng 1.","submitted_v1; event còn review; đúng lớp.","submitted_v2 / needs_revision_v1 / failed.",
   [pair("Chọn Approve/Revision/Reject và nhập ghi chú.","Kiểm tra quyền, event type, status và review_open."),
    pair("", "Nếu Approve, kiểm tra toàn bộ điều kiện hồ sơ."),
    pair("Xác nhận quyết định.","Ghi review round=1, cập nhật status/timeline và tạo notification.")],
   ["Revision→needs_revision_v1; sinh viên chỉnh sửa rồi nộp lại."],
   ["Hồ sơ chưa đủ điều kiện approve: thông báo phần chưa đạt và giữ submitted_v1."]),
mk("UC24","Theo dõi lớp được phân công","Cán bộ khoa","Không","Xem các lớp trong faculty_class_assignments.","Faculty officer context.","Dữ liệu lớp trong phạm vi được hiển thị.",
   simple_view("Lớp được phân công","Cán bộ khoa","dữ liệu lớp/hoạt động/hồ sơ","theo faculty_class_assignments")),
mk("UC25","Điểm danh QR/thủ công","Cán bộ khoa, Quản trị viên","Người tham gia","Xác nhận người tham gia thực tế.","Đủ quyền; activity published, đang diễn ra; đã đăng ký.","attended + log + ĐRL/notification.",
   [pair("Chọn activity và quét QR hoặc nhập MSSV.","Xác thực QR/MSSV, quyền và activity."),
    pair("", "Kiểm tra thời gian, user, registration và duplicate."),
    pair("Xác nhận điểm danh thủ công khi được yêu cầu.","Trong transaction, cập nhật attended, ghi log, conduct score và notification."),
    pair("", "Trả kết quả điểm danh.")],
   ["Có thể dùng MSSV thay QR."],
   ["invalid/out_of_window/not_registered/duplicate: ghi kết quả phù hợp và không cộng attendance lần nữa."]),
mk("UC26","Xem hồ sơ chờ duyệt vòng 2","Cán bộ khoa","Sinh viên","Xem queue submitted_v2 thuộc lớp phụ trách.","Faculty officer context; có assignment.","Queue vòng 2 hiển thị.",
   simple_view("Duyệt vòng 2","Cán bộ khoa","submission submitted_v2","theo faculty_class_assignments")),
mk("UC27","Đánh giá tiêu chí vòng 2","Cán bộ khoa","Sinh viên","Đánh giá tiêu chí thủ công ở vòng 2.","submitted_v2; đúng assignment; có evidence; không auto-pass.","criteria review round=2.",
   [pair("Mở tiêu chí cần đánh giá.","Kiểm tra evidence, auto result và assignment."),
    pair("Chọn Đạt/Chưa đạt và nhập ghi chú (nếu có).","Upsert submission_criteria_reviews round=2."),
    pair("", "Hiển thị trạng thái đánh giá mới.")],
   exc=["Sai assignment/auto-pass/thiếu evidence: từ chối cập nhật."]),
mk("UC28","Duyệt hồ sơ vòng 2","Cán bộ khoa","Sinh viên","Chốt kết quả cuối student submission.","submitted_v2; event còn review; đúng assignment.","passed / failed / needs_revision_v2 theo logic ứng dụng.",
   [pair("Chọn Approve/Revision/Reject và nhập ghi chú.","Kiểm tra quyền, assignment, status và review_open."),
    pair("", "Nếu Approve, kiểm tra toàn bộ điều kiện hồ sơ."),
    pair("Xác nhận quyết định.","Ghi review round=2, cập nhật status/timeline và notification.")],
   ["Revision→needs_revision_v2 theo logic ứng dụng."],
   ["Hồ sơ chưa đủ điều kiện approve hoặc schema chưa đồng bộ needs_revision_v2: không hoàn tất chuyển trạng thái."]),
mk("UC29","Xem hoạt động cán bộ","Cán bộ lớp, Cán bộ khoa","Không","Xem activity audience=officer.","Officer participant context.","Danh sách được hiển thị.",
   simple_view("Hoạt động cán bộ","Cán bộ","activity officer published","theo term/scope")),
mk("UC30","Đăng ký hoạt động cán bộ","Cán bộ lớp, Cán bộ khoa","Không","Đăng ký activity nội bộ dành cho cán bộ.","Activity officer published, đúng thời gian/scope, internal.","registered.",
   [pair("Chọn Đăng ký.","Kiểm tra audience, status, thời gian, source, lớp/scope và trùng."),
    pair("", "Tạo registration=registered và thông báo thành công.")],
   exc=["Không đủ điều kiện: hiển thị lý do và giữ nguyên dữ liệu."]),
mk("UC31","Xem đợt xét cán bộ","Cán bộ lớp, Cán bộ khoa","Không","Xem event type=officer.","Officer participant context.","Danh sách/chi tiết event hiển thị.",
   simple_view("Đợt xét cán bộ","Cán bộ","events type=officer","theo term")),
mk("UC32","Khởi tạo hồ sơ cán bộ","Cán bộ lớp, Cán bộ khoa","Không","Tạo/tiếp tục officer submission.","Event officer published và đang nhận.","Draft tồn tại, auto/score được tính.",
   [pair("Chọn Tham gia đợt xét.","Kiểm tra event phase và submission hiện có."),
    pair("", "Nếu chưa có, tạo submission=draft."),
    pair("", "Chạy auto-evaluation và evaluateSubmission; mở hồ sơ.")]),
mk("UC33","Bổ sung minh chứng cán bộ","Cán bộ lớp, Cán bộ khoa","Không","Lưu evidence và cập nhật score.","Submission thuộc user; draft/needs_revision_v1; còn hạn.","Evidence và score cập nhật.",
   [pair("Chọn tiêu chí cần bổ sung.","Hiển thị trạng thái auto/matched activity và vùng evidence."),
    pair("Nhập note/tệp và chọn Lưu.","Kiểm tra quyền/status/phase rồi lưu evidence."),
    pair("", "Chạy auto-evaluation và evaluateSubmission; hiển thị score mới.")],
   exc=["Hồ sơ không editable hoặc hết hạn: không lưu."]),
mk("UC34","Nộp hồ sơ cán bộ","Cán bộ lớp, Cán bộ khoa","Quản trị viên","Gửi officer submission vào queue Admin.","draft/needs_revision_v1; event đang nhận.","submitted_v1; score cập nhật.",
   [pair("Chọn Nộp hồ sơ.","Kiểm tra quyền, status và phase."),
    pair("", "Chạy auto-evaluation, requirements và evaluateSubmission."),
    pair("Xác nhận nộp.","Cập nhật submitted_v1/submitted_at và làm mới queue Admin.")],
   exc=["Thiếu requirements: hiển thị tiêu chí chưa đạt và quay lại chỉnh sửa."]),
mk("UC35","Theo dõi hồ sơ cán bộ","Cán bộ lớp, Cán bộ khoa","Quản trị viên","Theo dõi criteria, matched activities, score, evidence và review.","Có officer submission thuộc user.","Chi tiết hiển thị.",
   simple_view("Hồ sơ xét cán bộ","Cán bộ","officer submission, evidence, matched activities, score và review")),
mk("UC36","Quản lý học kỳ","Quản trị viên","Không","Tạo term và đặt term hiện hành.","Admin context.","Academic term cập nhật.",
   [pair("Mở Học kỳ.","Hiển thị danh sách term."),
    pair("Nhập thông tin kỳ và chọn Tạo.","Kiểm tra dữ liệu/thời gian và tạo term inactive."),
    pair("Chọn Đặt hiện hành cho một kỳ.","Đưa tất cả term về inactive rồi set term chọn is_active=1.")],
   exc=["Thời gian hoặc dữ liệu không hợp lệ: không tạo/cập nhật term."]),
mk("UC37","Quản lý người dùng","Quản trị viên","Không","Cập nhật role/trạng thái người dùng.","Admin context.","users/assignments liên quan cập nhật.",
   [pair("Mở chi tiết người dùng.","Hiển thị user, lớp và assignment."),
    pair("Chọn role/trạng thái và nhập lý do.","Kiểm tra rule của role/lớp."),
    pair("Xác nhận thay đổi.","Cập nhật user, assignment liên quan, audit/notification.")],
   exc=["Target/role/lớp không hợp lệ: hiển thị lỗi và không cập nhật."]),
mk("UC38","Quản lý lớp của người dùng","Quản trị viên","Người dùng","Chuyển lớp và bảo toàn lịch sử membership.","Admin; lớp đích active.","Membership cập nhật.",
   [pair("Chọn lớp mới và nhập lý do.","Kiểm tra user/lớp đích."),
    pair("Xác nhận chuyển lớp.","Đóng membership cũ, mở membership mới, xử lý role nếu cần và ghi audit.")],
   exc=["Lớp/user không hợp lệ: không thay đổi membership."]),
mk("UC39","Import danh sách sinh viên","Quản trị viên","Không","Import sinh viên với preview/validation.","Admin; có file và lớp đích.","users/class_members cập nhật.",
   [pair("Chọn lớp và tải file.","Parse file và phân loại từng dòng."),
    pair("Xem preview, chọn policy cập nhật/chuyển lớp.","Kiểm tra lựa chọn và các xung đột."),
    pair("Chọn Import.","Tạo/cập nhật user và class_members; trả thống kê kết quả.")],
   exc=["File/dòng/lớp không hợp lệ hoặc transfer chưa xác nhận: dừng import phần tương ứng."]),
mk("UC40","Phân công cán bộ khoa quản lý lớp","Quản trị viên","Cán bộ khoa","Gán tập lớp cho faculty officer.","Target faculty_officer; lớp active.","faculty_class_assignments cập nhật.",
   [pair("Chọn cán bộ khoa và các lớp phụ trách.","Kiểm tra target role và class ids."),
    pair("Chọn Lưu.","Thay thế tập assignments trong transaction và thông báo thành công.")],
   exc=["Target/lớp không hợp lệ: không cập nhật assignment."]),
mk("UC41","Quản lý hoạt động","Quản trị viên","Sinh viên/Cán bộ","Tạo/sửa/công bố/xóa activity.","Admin; có active term.","Activity cập nhật đúng rule.",
   [pair("Mở Quản lý hoạt động.","Hiển thị activity theo term."),
    pair("Nhập hoặc chỉnh sửa thông tin hoạt động.","Kiểm tra title, audience, organizer, thời gian, ĐRL, source và scope."),
    pair("Chọn Lưu/Công bố/Xóa.","Kiểm tra ràng buộc dữ liệu phát sinh và thực hiện thao tác hợp lệ.")],
   exc=["Ngoài term/trùng tên/điểm vượt max/có registration hoặc kỳ cũ: từ chối thao tác tương ứng."]),
mk("UC42","Khóa/mở đăng ký hoạt động","Quản trị viên","Sinh viên/Cán bộ","Điều khiển registration_locked.","Activity active term và published.","Cờ khóa cập nhật.",
   [pair("Chọn Khóa/Mở đăng ký.","Kiểm tra activity, term và status."),
    pair("Xác nhận.","Cập nhật registration_locked và làm mới giao diện.")],
   exc=["Activity kỳ cũ/chưa published: không cập nhật."]),
mk("UC43","Quản lý danh sách điểm danh","Quản trị viên, Cán bộ khoa","Người tham gia","Theo dõi và điều chỉnh attendance theo quyền.","Đủ quyền; activity tồn tại.","Attendance cập nhật.",
   [pair("Mở danh sách điểm danh của activity.","Đọc activity_registrations."),
    pair("Chọn thao tác điểm danh/điều chỉnh.","Kiểm tra quyền, user và trạng thái."),
    pair("Xác nhận thao tác.","Cập nhật attendance và làm mới danh sách.")],
   exc=["Không đủ quyền hoặc dữ liệu không hợp lệ: từ chối cập nhật."]),
mk("UC44","Import kết quả hoạt động","Quản trị viên","Người tham gia","Nạp attendance từ XLSX/XLS/CSV.","Activity đã kết thúc; file hợp lệ; có reason.","Attendance/ĐRL/import/audit/submission liên quan cập nhật.",
   [pair("Chọn activity, file, mode và nhập lý do.","Kiểm tra activity, thời gian, file và reason."),
    pair("", "Đọc MSSV, đối chiếu user/lớp/scope/registration và hiển thị preview/cảnh báo."),
    pair("Xác nhận tiếp tục.","Tạo import batch; xử lý từng row; upsert attended/conduct score; ghi row/audit."),
    pair("", "Re-evaluate auto criteria/score cho submission bị ảnh hưởng.")],
   ["Chọn append hoặc replace theo nhu cầu."],
   ["File/activity/user/scope không hợp lệ: ghi lỗi dòng hoặc dừng batch theo mức độ."]),
mk("UC45","Quản lý khung điểm rèn luyện","Quản trị viên","Sinh viên","Xây dựng conduct_score_categories phân cấp.","Admin; term tồn tại.","Cây category cập nhật.",
   [pair("Mở Khung điểm rèn luyện.","Hiển thị cây category."),
    pair("Thêm/sửa mục: code, tên, score_max, parent.","Kiểm tra trùng cùng cấp, code và parent."),
    pair("Chọn Lưu/Xóa.","Cập nhật cấu trúc khi thỏa ràng buộc.")],
   exc=["Code/tên/score/parent không hợp lệ: hiển thị lỗi và giữ dữ liệu cũ."]),
mk("UC46","Quản lý bộ tiêu chuẩn","Quản trị viên","Không","Tạo/clone/xóa criteria template.","Admin context.","Template/cấu trúc được quản lý.",
   [pair("Mở Bộ tiêu chuẩn.","Hiển thị criteria_templates."),
    pair("Chọn Tạo/Sao chép/Xóa.","Kiểm tra for_type và trạng thái đã sử dụng."),
    pair("Xác nhận.","Tạo/clone cấu trúc hoặc xóa khi mutable.")],
   ["Template đã dùng: chọn clone để tạo phiên bản mới."],
   ["Template đã được event tham chiếu: không cho sửa/xóa trực tiếp."]),
mk("UC47","Cấu hình tiêu chuẩn và tiêu chí","Quản trị viên","Không","Cấu hình group/item trong template.","Template mutable.","Groups/items cập nhật.",
   [pair("Mở Cấu hình template.","Kiểm tra template mutable."),
    pair("Nhập thông tin tiêu chuẩn/tiêu chí.","Kiểm tra min_required, score_max, evidence và group."),
    pair("Chọn Lưu.","Sinh code/sort_order và lưu group/item.")],
   exc=["Template đã dùng hoặc dữ liệu không hợp lệ: không lưu."]),
mk("UC48","Gắn hoạt động vào tiêu chí","Quản trị viên","Không","Tạo activity rule và score_value.","Template mutable; criteria/activity hợp lệ.","Rule cập nhật.",
   [pair("Chọn tiêu chí, activity và nhập score_value.","Kiểm tra template, criteria, activity và score."),
    pair("Chọn Lưu.","Upsert criteria_activity_rules và làm mới cấu hình.")],
   exc=["Criteria/activity/score không hợp lệ hoặc template locked: không lưu."]),
mk("UC49","Gắn điều kiện ĐRL vào tiêu chí","Quản trị viên","Không","Gắn ngưỡng conduct score.","Template mutable; criteria hợp lệ.","Conduct rule cập nhật.",
   [pair("Nhập min_score và period_scope.","Kiểm tra min_score và scope current/any."),
    pair("Chọn Lưu.","Upsert criteria_conduct_rules.")],
   exc=["Giá trị không hợp lệ hoặc template locked: không lưu."]),
mk("UC50","Quản lý đợt xét","Quản trị viên","Sinh viên/Cán bộ","Tạo/sửa/công bố/đóng/xóa event.","Admin; active term; template phù hợp.","Event cập nhật.",
   [pair("Mở Quản lý đợt xét.","Hiển thị events theo term."),
    pair("Nhập/chỉnh sửa title, type, template, time, allow_late, status.","Kiểm tra type-template, thời gian và trạng thái sử dụng."),
    pair("Chọn Lưu/Công bố/Đóng/Xóa.","Thực hiện thao tác nếu không vi phạm dữ liệu lịch sử.")],
   exc=["Template sai type, thời gian sai, event kỳ cũ hoặc đã có submission: từ chối thao tác không hợp lệ."]),
mk("UC51","Đồng bộ snapshot tiêu chuẩn","Quản trị viên","Không","Copy template groups/items sang event snapshot.","Event có template; chưa có submission.","Snapshot cập nhật.",
   [pair("Chọn Đồng bộ tiêu chuẩn.","Kiểm tra event/template và đếm submission."),
    pair("Xác nhận đồng bộ.","Xóa snapshot cũ và copy groups/items từ template."),
    pair("", "Thông báo đồng bộ thành công.")],
   exc=["Event đã có submission: không đồng bộ để bảo toàn tiêu chí lịch sử."]),
mk("UC52","Theo dõi hồ sơ sinh viên","Quản trị viên","Sinh viên/Reviewer","Giám sát student submissions.","Admin context.","Danh sách/chi tiết hiển thị.",
   simple_view("Hồ sơ sinh viên","Quản trị viên","student submissions, evidence, auto results, reviews và timeline","theo term/event/lớp/status")),
mk("UC53","Theo dõi hồ sơ cán bộ","Quản trị viên","Cán bộ","Giám sát officer submissions và score.","Admin context.","Danh sách/chi tiết hiển thị.",
   simple_view("Hồ sơ cán bộ","Quản trị viên","officer submissions, matched activities, evidence, score và review","theo term/event/status")),
mk("UC54","Duyệt hồ sơ cán bộ","Quản trị viên","Cán bộ lớp/Cán bộ khoa","Admin duyệt officer submission một vòng.","Officer submission=submitted_v1.","passed/failed/needs_revision_v1.",
   [pair("Mở hồ sơ và chọn pass/fail/revise; nhập ghi chú.","Kiểm tra event officer và status."),
    pair("", "Nếu pass, chạy auto-evaluation và kiểm tra approval requirements."),
    pair("Xác nhận quyết định.","Ghi review round=1, cập nhật status, notification và tính lại score.")],
   ["revise→needs_revision_v1; cán bộ chỉnh sửa rồi nộp lại."],
   ["Không phải officer/sai status/không đủ điều kiện pass: từ chối quyết định."]),
mk("UC55","Xem dashboard và báo cáo","Quản trị viên","Không","Xem KPI, biểu đồ và bảng báo cáo theo học kỳ.","Admin context; term hợp lệ.","Báo cáo hiển thị.",
   [pair("Mở Dashboard/Báo cáo.","Xác định selectedTerm và tổng hợp dữ liệu."),
    pair("Chọn report và bộ lọc.","Áp dụng filter, tính KPI/tỷ lệ/bảng/chart."),
    pair("", "Hiển thị kết quả.")]),
mk("UC56","Xuất báo cáo","Quản trị viên","Không","Xuất báo cáo theo bộ lọc hiện tại.","Admin context.","Tệp export được trả về.",
   [pair("Chọn Xuất báo cáo.","Đọc term/event/class/top từ bộ lọc."),
    pair("", "Truy vấn dữ liệu cùng phạm vi và tạo tệp."),
    pair("Chọn lưu tệp trên trình duyệt.","Trả tệp export.")],
   exc=["Không có dữ liệu phù hợp: thông báo hoặc xuất tệp không có dòng dữ liệu."]),
mk("UC57","Xem nhật ký hệ thống","Quản trị viên","Không","Tra cứu audit_logs.","Admin context.","Danh sách audit hiển thị.",
   [pair("Mở Nhật ký hệ thống.","Truy vấn audit_logs và actor/entity liên quan."),
    pair("Chọn bộ lọc hoặc một bản ghi.","Hiển thị action, time, reason và before/after nếu có.")])
]

# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------
def clean_actor(text):
    t = (text or "").strip()
    for p in ["Người dùng ", "Sinh viên ", "Cán bộ lớp ", "Cán bộ khoa ", "Quản trị viên ", "Admin "]:
        if t.lower().startswith(p.lower()):
            t = t[len(p):]
            break
    return t[:1].upper() + t[1:] if t else t

def clean_system(text):
    t = (text or "").strip()
    for p in ["Hệ thống ", "System "]:
        if t.lower().startswith(p.lower()):
            t = t[len(p):]
            break
    return t[:1].upper() + t[1:] if t else t

def flatten_main(u):
    steps = []
    n = 1
    for actor, system in u["main"]:
        if actor and actor.strip():
            steps.append((n, "Actor", humanize(clean_actor(actor))))
            n += 1
        if system and system.strip():
            steps.append((n, "Hệ thống", humanize(clean_system(system))))
            n += 1
    return steps

def branch_base(main_steps):
    # Prefer a system validation/decision step because this is where most alternate/exception branches originate.
    candidates = []
    for n, side, action in main_steps:
        if side == "Hệ thống":
            low = action.lower()
            if any(k in low for k in ["kiểm tra", "xác thực", "đối chiếu", "tính ", "lọc ", "tìm ", "truy vấn", "đọc "]):
                candidates.append(n)
    if candidates:
        return candidates[0]
    return 2 if len(main_steps) >= 2 else 1

ACTOR_VERBS = (
    "chọn","nhập","xác nhận","mở","nộp","đăng ký","hủy","sửa","tải","giữ","đóng",
    "tiếp tục","thực hiện","quét","nhấn","chuyển","xem","yêu cầu","bổ sung"
)
def side_for_branch(text):
    low = text.strip().lower()
    return "Actor" if low.startswith(ACTOR_VERBS) else "Hệ thống"

def branch_actions(items, main_steps, kind):
    if not items or items == ["Không có."]:
        return []
    base = branch_base(main_steps)
    rows = []
    sub = 1
    for raw in items:
        raw = str(raw).strip()
        if not raw or raw.lower() == "không có.":
            continue
        side = side_for_branch(raw)
        txt = humanize(clean_actor(raw) if side == "Actor" else clean_system(raw))
        # Rewrite descriptive conditions into action-oriented language for the System column.
        if side == "Hệ thống" and not re.match(r"^(Kiểm tra|Hiển thị|Từ chối|Cho phép|Giữ|Áp dụng|Chuyển|Cập nhật|Ghi|Trả|Kết thúc|Phát hiện|Xác định|Yêu cầu|Thông báo|Không|Đưa|Khôi phục|Chỉ|Sử dụng|Thực hiện)", txt):
            txt = "Xử lý trường hợp: " + txt[:1].lower() + txt[1:]
        rows.append((f"{base}.{sub}", side, txt))
        sub += 1

    if rows:
        joined = " ".join(str(x) for x in items).lower()
        recoverable = any(k in joined for k in ["sai", "thiếu", "không hợp lệ", "trùng", "vượt", "nhập lại", "quay lại", "chọn lại"])
        if kind == "alternative" or recoverable:
            back = max(1, base - 1)
            rows.append((f"{base}.{sub}", "Actor", f"Chọn quay lại bước {back} để điều chỉnh hoặc thực hiện lựa chọn khác."))
        else:
            rows.append((f"{base}.{sub}", "Hệ thống", "Kết thúc use case và không thực hiện thay đổi dữ liệu ngoài các ghi nhận lỗi/audit cần thiết."))
    return rows

def uc_role(u):
    a = u["actor"]
    roles = []
    if "Sinh viên" in a: roles.append("Sinh viên")
    if "Cán bộ lớp" in a: roles.append("Cán bộ lớp")
    if "Cán bộ khoa" in a: roles.append("Cán bộ khoa")
    if "Quản trị viên" in a: roles.append("Quản trị viên")
    if a == "Người dùng":
        roles.extend(["Sinh viên","Cán bộ lớp","Cán bộ khoa","Quản trị viên"])
    return roles


def humanize(text):
    """Chuyển thuật ngữ triển khai thành ngôn ngữ nghiệp vụ dễ hiểu."""
    s = str(text)
    replacements = [
        ("Officer participant context", "giao diện dành cho cán bộ"),
        ("Faculty officer context", "giao diện cán bộ khoa"),
        ("Class officer context", "giao diện cán bộ lớp"),
        ("Student context", "giao diện sinh viên"),
        ("Admin context", "giao diện quản trị"),
        ("login_context", "giao diện theo vai trò"),
        ("activity_registrations", "thông tin đăng ký tham gia"),
        ("registration_locked", "trạng thái khóa đăng ký"),
        ("conduct_score_categories", "khung điểm rèn luyện"),
        ("conduct_scores", "điểm rèn luyện"),
        ("criteria_template_groups/items", "nhóm tiêu chuẩn và tiêu chí"),
        ("criteria_activity_rules", "quy tắc ghi nhận từ hoạt động"),
        ("criteria_conduct_rules", "điều kiện điểm rèn luyện"),
        ("event_criteria_groups/items", "bộ tiêu chuẩn áp dụng cho đợt xét"),
        ("faculty_class_assignments", "phân công lớp cho cán bộ khoa"),
        ("submission_items/files", "minh chứng của hồ sơ"),
        ("submission_criteria_reviews", "kết quả đánh giá từng tiêu chí"),
        ("submission_timeline", "lịch sử xử lý hồ sơ"),
        ("notifications", "thông báo"),
        ("audit_logs", "nhật ký hệ thống"),
        ("activity", "hoạt động"),
        ("Activity", "Hoạt động"),
        ("events", "các đợt xét"),
        ("event", "đợt xét"),
        ("Event", "Đợt xét"),
        ("submission", "hồ sơ"),
        ("Submission", "Hồ sơ"),
        ("registration", "đăng ký tham gia"),
        ("Registration", "Đăng ký tham gia"),
        ("published", "đã công bố"),
        ("scope", "phạm vi áp dụng"),
        ("role", "vai trò"),
        ("context", "giao diện theo vai trò"),
        ("auto-evaluation", "đánh giá tự động"),
        ("auto evaluation", "đánh giá tự động"),
        ("auto-evaluate", "đánh giá tự động"),
        ("score_total", "tổng điểm"),
        ("score_max", "điểm tối đa"),
        ("evidence", "minh chứng"),
        ("review", "xét duyệt"),
        ("timeline", "lịch sử xử lý"),
        ("queue", "danh sách chờ xử lý"),
        ("attended", "đã tham gia"),
        ("student", "sinh viên"),
        ("officer", "cán bộ"),
        ("Admin", "Quản trị viên"),
    ]
    for old,new in replacements:
        s=s.replace(old,new)
    s=re.sub(r"\bstatus\b","trạng thái",s,flags=re.I)
    s=re.sub(r"\bphase\b","giai đoạn",s,flags=re.I)
    s=re.sub(r"\bterm\b","học kỳ",s,flags=re.I)
    s=re.sub(r"\bsource\b","nguồn ghi nhận",s,flags=re.I)
    s=re.sub(r"\bfile\b","tệp",s,flags=re.I)
    s=re.sub(r"\bnote\b","ghi chú",s,flags=re.I)
    return s.replace("_"," ")

# ------------------------------------------------------------
# Diagrams
# ------------------------------------------------------------
def savefig(path):
    plt.savefig(path, dpi=180, bbox_inches="tight")
    plt.close()

def draw_actor(ax, x, y, label):
    ax.add_patch(Circle((x, y+0.52), 0.16, fill=False, linewidth=1.1))
    ax.plot([x,x],[y+0.36,y-0.12],lw=1.1)
    ax.plot([x-0.25,x+0.25],[y+0.16,y+0.16],lw=1.1)
    ax.plot([x,x-0.23],[y-0.12,y-0.50],lw=1.1)
    ax.plot([x,x+0.23],[y-0.12,y-0.50],lw=1.1)
    ax.text(x,y-0.70,label,ha="center",va="top",fontsize=8)

def usecase_image(filename, title, actor, labels):
    cols = 3 if len(labels)>12 else 2
    rows = math.ceil(len(labels)/cols)
    fig, ax = plt.subplots(figsize=(13, max(7, rows*1.4)))
    ax.axis("off")
    ax.set_xlim(0, 14); ax.set_ylim(0, max(10, rows*1.6+2))
    H=max(10, rows*1.6+2)
    draw_actor(ax,1.1,H/2,actor)
    ax.add_patch(Rectangle((2.7,0.45),10.8,H-0.9,fill=False,lw=1.2))
    ax.text(8.1,H-0.8,title,ha="center",fontsize=12,fontweight="bold")
    ytop=H-1.8
    for i,label in enumerate(labels):
        c=i%cols; r=i//cols
        x=4.4+c*3.15; y=ytop-r*1.25
        ax.add_patch(Ellipse((x,y),2.55,0.75,fill=False,lw=0.9))
        ax.text(x,y,"\n".join(textwrap.wrap(label,22)),ha="center",va="center",fontsize=7)
        ax.plot([1.4,x-1.3],[H/2,y],lw=0.45)
    path=IMG/filename
    savefig(path)
    return path

def overall_usecase_image():
    fig, ax = plt.subplots(figsize=(14,10))
    ax.axis("off"); ax.set_xlim(0,15); ax.set_ylim(0,11)
    ax.add_patch(Rectangle((3.0,0.5),11.2,10,fill=False,lw=1.2))
    ax.text(8.6,10.1,"MeritTrack – Use Case tổng quát",ha="center",fontsize=13,fontweight="bold")
    actor_pos={"Sinh viên":(1.1,8.4),"Cán bộ lớp":(1.1,6.2),"Cán bộ khoa":(1.1,4.0),"Quản trị viên":(1.1,1.8)}
    for a,(x,y) in actor_pos.items(): draw_actor(ax,x,y,a)
    labels=[
      ("Xác thực & tài khoản",5.0,8.8),("Tham gia hoạt động",8.0,8.8),("Điểm rèn luyện",11.0,8.8),
      ("Hồ sơ xét sinh viên",5.0,7.1),("Duyệt vòng 1",8.0,7.1),("Duyệt vòng 2",11.0,7.1),
      ("Nghiệp vụ cán bộ",5.0,5.4),("Quản lý học kỳ",8.0,5.4),("Người dùng & lớp",11.0,5.4),
      ("Hoạt động & điểm danh",5.0,3.7),("Khung điểm ĐRL",8.0,3.7),("Bộ tiêu chuẩn",11.0,3.7),
      ("Đợt xét",5.0,2.0),("Duyệt hồ sơ cán bộ",8.0,2.0),("Báo cáo & nhật ký",11.0,2.0)
    ]
    ucpos={}
    for label,x,y in labels:
        ucpos[label]=(x,y)
        ax.add_patch(Ellipse((x,y),2.45,0.72,fill=False,lw=0.9))
        ax.text(x,y,"\n".join(textwrap.wrap(label,20)),ha="center",va="center",fontsize=7)
    links={
      "Sinh viên":["Xác thực & tài khoản","Tham gia hoạt động","Điểm rèn luyện","Hồ sơ xét sinh viên"],
      "Cán bộ lớp":["Xác thực & tài khoản","Duyệt vòng 1","Nghiệp vụ cán bộ"],
      "Cán bộ khoa":["Xác thực & tài khoản","Duyệt vòng 2","Nghiệp vụ cán bộ","Hoạt động & điểm danh"],
      "Quản trị viên":["Quản lý học kỳ","Người dùng & lớp","Hoạt động & điểm danh","Khung điểm ĐRL","Bộ tiêu chuẩn","Đợt xét","Duyệt hồ sơ cán bộ","Báo cáo & nhật ký"]
    }
    for a,ucs in links.items():
        x1,y1=actor_pos[a]
        for u in ucs:
            x2,y2=ucpos[u]
            ax.plot([x1+0.3,x2-1.25],[y1,y2],lw=0.45)
    path=IMG/"uc_overall.png"; savefig(path); return path

def swimlane(filename, title, lanes, steps):
    # steps: (lane_index, label, kind, target_index_optional)
    n=len(steps); W=12; H=max(7, n*0.75+2)
    fig,ax=plt.subplots(figsize=(12,H*0.65))
    ax.axis("off"); ax.set_xlim(0,W); ax.set_ylim(0,H)
    lane_w=W/len(lanes)
    for i,l in enumerate(lanes):
        x=i*lane_w
        ax.add_patch(Rectangle((x,0.3),lane_w,H-0.8,fill=False,lw=0.9))
        ax.text(x+lane_w/2,H-0.65,l,ha="center",fontsize=9,fontweight="bold")
    ypos=[]
    for i,s in enumerate(steps):
        y=H-1.4-i*((H-2.1)/max(1,n-1))
        ypos.append(y)
        lane,label,kind=s[:3]
        x=lane*lane_w+lane_w/2
        if kind=="decision":
            ax.add_patch(Polygon([[x,y+0.30],[x+0.75,y],[x,y-0.30],[x-0.75,y]],closed=True,fill=False,lw=0.9))
        elif kind=="start" or kind=="end":
            ax.add_patch(Ellipse((x,y),1.5,0.45,fill=False,lw=0.9))
        else:
            ax.add_patch(Rectangle((x-0.9,y-0.28),1.8,0.56,fill=False,lw=0.9))
        ax.text(x,y,"\n".join(textwrap.wrap(label,28)),ha="center",va="center",fontsize=7)
        if i>0:
            pl=steps[i-1][0]
            px=pl*lane_w+lane_w/2
            py=ypos[i-1]
            ax.add_patch(FancyArrowPatch((px,py-0.3),(x,y+0.3),arrowstyle="->",mutation_scale=9,lw=0.7))
    ax.text(W/2,H-0.15,title,ha="center",fontsize=12,fontweight="bold")
    path=IMG/filename; savefig(path); return path

OVERALL=overall_usecase_image()
student_labels=[f'{u["id"]} {u["name"]}' for u in USECASES if "Sinh viên" in uc_role(u) and int(u["id"][2:])<=20]
class_labels=[f'{u["id"]} {u["name"]}' for u in USECASES if "Cán bộ lớp" in uc_role(u) and (int(u["id"][2:])<=8 or 21<=int(u["id"][2:])<=23 or 29<=int(u["id"][2:])<=35)]
faculty_labels=[f'{u["id"]} {u["name"]}' for u in USECASES if "Cán bộ khoa" in uc_role(u) and (int(u["id"][2:])<=8 or 24<=int(u["id"][2:])<=35)]
admin_labels=[f'{u["id"]} {u["name"]}' for u in USECASES if "Quản trị viên" in uc_role(u) and (int(u["id"][2:])<=8 or int(u["id"][2:])>=36)]
UC_STUDENT=usecase_image("uc_student.png","Phân hệ Sinh viên","Sinh viên",student_labels)
UC_CLASS=usecase_image("uc_class_officer.png","Phân hệ Cán bộ lớp","Cán bộ lớp",class_labels)
UC_FACULTY=usecase_image("uc_faculty_officer.png","Phân hệ Cán bộ khoa","Cán bộ khoa",faculty_labels)
UC_ADMIN=usecase_image("uc_admin.png","Phân hệ Quản trị viên","Quản trị viên",admin_labels)

BUSINESS_DIAGRAMS = [
    ("BP01", swimlane("bp01_auth.png","BP01 – Đăng nhập và chuyển ngữ cảnh",
        ["Người dùng","MeritTrack"],
        [(0,"Bắt đầu","start"),(0,"Nhập MSSV/email và mật khẩu","task"),(1,"Kiểm tra tài khoản và mật khẩu","decision"),
         (1,"Tạo session, đặt context mặc định","task"),(0,"Chọn chuyển context (nếu cần)","task"),(1,"Kiểm tra role-context","decision"),
         (1,"Cập nhật session và dashboard","task"),(0,"Kết thúc","end")])),
    ("BP02", swimlane("bp02_activity.png","BP02 – Đăng ký và điểm danh hoạt động",
        ["Người tham gia","MeritTrack","Cán bộ khoa/Admin"],
        [(0,"Xem hoạt động","start"),(1,"Hiển thị activity phù hợp","task"),(0,"Chọn Đăng ký","task"),(1,"Kiểm tra điều kiện đăng ký","decision"),
         (1,"Ghi registration=registered","task"),(0,"Tham gia hoạt động","task"),(2,"Quét QR / nhập MSSV","task"),
         (1,"Kiểm tra thời gian & registration","decision"),(1,"Cập nhật attended + log + ĐRL","task"),(0,"Nhận thông báo","end")])),
    ("BP03", swimlane("bp03_student_submission.png","BP03 – Xét hồ sơ sinh viên hai vòng",
        ["Sinh viên","MeritTrack","Cán bộ lớp","Cán bộ khoa"],
        [(0,"Tạo/tiếp tục hồ sơ","start"),(1,"Auto-evaluate","task"),(0,"Bổ sung minh chứng","task"),(0,"Nộp hồ sơ","task"),
         (1,"Kiểm tra requirements → submitted_v1","decision"),(2,"Đánh giá tiêu chí và duyệt vòng 1","task"),(1,"submitted_v2 / revise / failed","decision"),
         (3,"Đánh giá tiêu chí và duyệt vòng 2","task"),(1,"passed / revise / failed","decision"),(0,"Theo dõi kết quả","end")])),
    ("BP04", swimlane("bp04_officer_submission.png","BP04 – Xét hồ sơ cán bộ",
        ["Cán bộ","MeritTrack","Admin"],
        [(0,"Tạo hồ sơ officer","start"),(1,"Auto-evaluate và tính điểm","task"),(0,"Bổ sung minh chứng","task"),(0,"Nộp hồ sơ","task"),
         (1,"Kiểm tra requirements → submitted_v1","decision"),(2,"Đánh giá hồ sơ","task"),(1,"passed / failed / needs_revision_v1","decision"),(0,"Theo dõi kết quả","end")])),
    ("BP05", swimlane("bp05_criteria_event.png","BP05 – Cấu hình bộ tiêu chuẩn và đợt xét",
        ["Admin","MeritTrack"],
        [(0,"Tạo/clone bộ tiêu chuẩn","start"),(1,"Kiểm tra template mutable","decision"),(0,"Cấu hình nhóm/tiêu chí/rule","task"),
         (1,"Lưu template","task"),(0,"Tạo đợt xét","task"),(1,"Kiểm tra type/template/term","decision"),
         (0,"Đồng bộ snapshot","task"),(1,"Copy event_criteria groups/items","task"),(0,"Công bố đợt xét","end")])),
    ("BP06", swimlane("bp06_import_report.png","BP06 – Import kết quả hoạt động và đồng bộ hồ sơ",
        ["Admin","MeritTrack","Hồ sơ liên quan"],
        [(0,"Chọn activity + file + mode + reason","start"),(1,"Parse và đối chiếu MSSV/lớp/scope","decision"),
         (0,"Xác nhận các trường hợp đặc biệt","task"),(1,"Ghi import batch/rows + attended + ĐRL","task"),
         (2,"Re-evaluate auto criteria và score","task"),(1,"Ghi audit và cập nhật báo cáo","task"),(0,"Kết thúc","end")]))
]

# ------------------------------------------------------------
# XMI 2.1 for Visual Paradigm
# ------------------------------------------------------------
def xml_escape(s):
    return str(s).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace('"',"&quot;")
def sid(s):
    s=re.sub(r"[^A-Za-z0-9_]+","_",s)
    return "_"+s

def make_xmi(path, model_name, actors, usecases, assocs, includes=None, extends=None):
    includes = includes or []
    extends = extends or []
    parts=['<?xml version="1.0" encoding="UTF-8"?>',
           '<xmi:XMI xmi:version="2.1" xmlns:xmi="http://schema.omg.org/spec/XMI/2.1" xmlns:uml="http://www.eclipse.org/uml2/3.0.0/UML">',
           f'  <uml:Model xmi:id="{sid(model_name)}" name="{xml_escape(model_name)}">']
    for a in actors:
        parts.append(f'    <packagedElement xmi:type="uml:Actor" xmi:id="{sid("actor_"+a)}" name="{xml_escape(a)}"/>')
    for u in usecases:
        parts.append(f'    <packagedElement xmi:type="uml:UseCase" xmi:id="{sid("uc_"+u)}" name="{xml_escape(u)}">')
        for i,(src,dst) in enumerate(includes,1):
            if src==u:
                parts.append(f'      <include xmi:type="uml:Include" xmi:id="{sid("inc_"+str(i)+"_"+src+"_"+dst)}" addition="{sid("uc_"+dst)}"/>')
        for i,(src,dst) in enumerate(extends,1):
            if src==u:
                parts.append(f'      <extend xmi:type="uml:Extend" xmi:id="{sid("ext_"+str(i)+"_"+src+"_"+dst)}" extendedCase="{sid("uc_"+dst)}"/>')
        parts.append('    </packagedElement>')
    for i,(a,u) in enumerate(assocs,1):
        aid=sid(f"assoc_{i}_{a}_{u}"); e1=sid(f"aend_{i}"); e2=sid(f"uend_{i}")
        parts.append(f'    <packagedElement xmi:type="uml:Association" xmi:id="{aid}" memberEnd="{e1} {e2}">')
        parts.append(f'      <ownedEnd xmi:type="uml:Property" xmi:id="{e1}" type="{sid("actor_"+a)}" association="{aid}"/>')
        parts.append(f'      <ownedEnd xmi:type="uml:Property" xmi:id="{e2}" type="{sid("uc_"+u)}" association="{aid}"/>')
        parts.append('    </packagedElement>')
    parts.extend(['  </uml:Model>','</xmi:XMI>'])
    path.write_text("\n".join(parts), encoding="utf-8")

def usecase_name(u):
    return f'{u["id"]} - {u["name"]}'

SUPPORT_UCS = [
    "Kiểm tra điều kiện đăng ký",
    "Đánh giá hồ sơ tự động",
    "Kiểm tra điều kiện nộp hồ sơ",
    "Gửi thông báo",
    "Kiểm tra quyền xử lý hồ sơ",
]
INCLUDE_ID_PAIRS = [
    ("UC11","Kiểm tra điều kiện đăng ký"),("UC30","Kiểm tra điều kiện đăng ký"),
    ("UC16","Đánh giá hồ sơ tự động"),("UC19","Đánh giá hồ sơ tự động"),
    ("UC32","Đánh giá hồ sơ tự động"),("UC34","Đánh giá hồ sơ tự động"),("UC54","Đánh giá hồ sơ tự động"),
    ("UC19","Kiểm tra điều kiện nộp hồ sơ"),("UC34","Kiểm tra điều kiện nộp hồ sơ"),
    ("UC23","Kiểm tra quyền xử lý hồ sơ"),("UC28","Kiểm tra quyền xử lý hồ sơ"),("UC54","Kiểm tra quyền xử lý hồ sơ"),
    ("UC23","Gửi thông báo"),("UC25","Gửi thông báo"),("UC28","Gửi thông báo"),("UC54","Gửi thông báo"),
]
EXTEND_ID_PAIRS = [
    ("UC11","UC10"),("UC12","UC10"),
    ("UC16","UC15"),
    ("UC17","UC20"),("UC18","UC20"),
    ("UC30","UC29"),
    ("UC32","UC31"),("UC33","UC35"),
]

by_id={u["id"]:u for u in USECASES}
def relation_names(selected):
    ids={u["id"] for u in selected}
    includes=[]
    extends=[]
    needed_support=set()
    for src,dst in INCLUDE_ID_PAIRS:
        if src in ids:
            includes.append((usecase_name(by_id[src]), dst))
            needed_support.add(dst)
    for src,dst in EXTEND_ID_PAIRS:
        if src in ids and dst in ids:
            extends.append((usecase_name(by_id[src]), usecase_name(by_id[dst])))
    return includes,extends,sorted(needed_support)

overall_actors=["Sinh viên","Cán bộ lớp","Cán bộ khoa","Quản trị viên"]
overall_ucs=["Đăng nhập và quản lý tài khoản","Tham gia hoạt động","Theo dõi điểm rèn luyện","Quản lý hồ sơ xét sinh viên","Duyệt hồ sơ vòng 1","Duyệt hồ sơ vòng 2","Tham gia nghiệp vụ cán bộ","Quản lý học kỳ","Quản lý người dùng và lớp","Quản lý hoạt động và điểm danh","Quản lý khung điểm rèn luyện","Quản lý bộ tiêu chuẩn","Quản lý đợt xét","Duyệt hồ sơ cán bộ","Xem báo cáo và nhật ký"]
overall_assoc=[]
omap={
"Sinh viên":["Đăng nhập và quản lý tài khoản","Tham gia hoạt động","Theo dõi điểm rèn luyện","Quản lý hồ sơ xét sinh viên"],
"Cán bộ lớp":["Đăng nhập và quản lý tài khoản","Duyệt hồ sơ vòng 1","Tham gia nghiệp vụ cán bộ"],
"Cán bộ khoa":["Đăng nhập và quản lý tài khoản","Duyệt hồ sơ vòng 2","Tham gia nghiệp vụ cán bộ","Quản lý hoạt động và điểm danh"],
"Quản trị viên":["Quản lý học kỳ","Quản lý người dùng và lớp","Quản lý hoạt động và điểm danh","Quản lý khung điểm rèn luyện","Quản lý bộ tiêu chuẩn","Quản lý đợt xét","Duyệt hồ sơ cán bộ","Xem báo cáo và nhật ký"]
}
for a,us in omap.items():
    overall_assoc += [(a,u) for u in us]
make_xmi(XMI/"01_MeritTrack_UseCase_TongQuan.xmi","MeritTrack - Use Case tổng quát",overall_actors,overall_ucs,overall_assoc)

def role_xmi(filename, actor, selected):
    base=[usecase_name(u) for u in selected]
    includes,extends,support=relation_names(selected)
    names=base+support
    make_xmi(XMI/filename, f"MeritTrack - {actor}", [actor], names, [(actor,n) for n in base], includes, extends)

student_selected=[u for u in USECASES if "Sinh viên" in uc_role(u) and int(u["id"][2:])<=20]
class_selected=[u for u in USECASES if "Cán bộ lớp" in uc_role(u) and (int(u["id"][2:])<=8 or 21<=int(u["id"][2:])<=23 or 29<=int(u["id"][2:])<=35)]
faculty_selected=[u for u in USECASES if "Cán bộ khoa" in uc_role(u) and (int(u["id"][2:])<=8 or 24<=int(u["id"][2:])<=35)]
admin_selected=[u for u in USECASES if "Quản trị viên" in uc_role(u) and (u["id"] in ["UC01","UC02","UC04","UC07","UC08"] or int(u["id"][2:])>=36)]
role_xmi("02_MeritTrack_UseCase_SinhVien.xmi","Sinh viên",student_selected)
role_xmi("03_MeritTrack_UseCase_CanBoLop.xmi","Cán bộ lớp",class_selected)
role_xmi("04_MeritTrack_UseCase_CanBoKhoa.xmi","Cán bộ khoa",faculty_selected)
role_xmi("05_MeritTrack_UseCase_Admin.xmi","Quản trị viên",admin_selected)

# Editable draw.io helpers ------------------------------------------------------
def drawio_doc(title, lanes, nodes, edges, path):
    # nodes: [{id,label,lane,kind,x,y,w,h}]
    # edges: [(src,dst,label,style)]
    cells=['<mxCell id="0"/>','<mxCell id="1" parent="0"/>']
    lane_ids={}
    lane_w=320
    total_h=max([n.get("y",80)+n.get("h",60) for n in nodes]+[600])+100
    for i,lane in enumerate(lanes):
        lid=f"lane{i+1}"; lane_ids[i]=lid
        x=20+i*lane_w
        cells.append(f'<mxCell id="{lid}" value="{xml_escape(lane)}" style="swimlane;horizontal=0;startSize=32;html=1;" vertex="1" parent="1"><mxGeometry x="{x}" y="20" width="{lane_w}" height="{total_h}" as="geometry"/></mxCell>')
    for n in nodes:
        kind=n.get("kind","task")
        style="rounded=1;whiteSpace=wrap;html=1;"
        if kind=="decision": style="rhombus;whiteSpace=wrap;html=1;"
        elif kind in ("start","end"): style="ellipse;whiteSpace=wrap;html=1;"
        elif kind=="actor": style="shape=umlActor;verticalLabelPosition=bottom;verticalAlign=top;html=1;"
        parent=lane_ids.get(n.get("lane",0),"1")
        cells.append(f'<mxCell id="{n["id"]}" value="{xml_escape(n["label"])}" style="{style}" vertex="1" parent="{parent}"><mxGeometry x="{n.get("x",60)}" y="{n.get("y",60)}" width="{n.get("w",190)}" height="{n.get("h",55)}" as="geometry"/></mxCell>')
    for i,e in enumerate(edges,1):
        src,dst,label,*rest=e
        style=rest[0] if rest else "edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=block;"
        cells.append(f'<mxCell id="e{i}" value="{xml_escape(label)}" style="{style}" edge="1" parent="1" source="{src}" target="{dst}"><mxGeometry relative="1" as="geometry"/></mxCell>')
    xml='<?xml version="1.0" encoding="UTF-8"?><mxfile host="app.diagrams.net"><diagram name="'+xml_escape(title)+'"><mxGraphModel dx="1422" dy="794" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="1169" pageHeight="827"><root>'+''.join(cells)+'</root></mxGraphModel></diagram></mxfile>'
    path.write_text(xml,encoding="utf-8")

def business_drawio(code,title,lanes,steps,path):
    nodes=[]; edges=[]
    for i,s in enumerate(steps):
        lane,label,kind=s[:3]
        nodes.append({"id":f"n{i+1}","label":label,"lane":lane,"kind":kind,"x":65,"y":60+i*85,"w":190,"h":55})
        if i>0: edges.append((f"n{i}",f"n{i+1}",""))
    drawio_doc(f"{code} - {title}",lanes,nodes,edges,path)

BUSINESS_EDITABLE = [
("BP01","Đăng nhập và chuyển ngữ cảnh",["Người dùng","Hệ thống"],[(0,"Bắt đầu","start"),(0,"Nhập thông tin đăng nhập","task"),(1,"Kiểm tra tài khoản","decision"),(1,"Tạo phiên đăng nhập","task"),(0,"Chọn giao diện theo vai trò (nếu cần)","task"),(1,"Kiểm tra quyền chuyển giao diện","decision"),(1,"Hiển thị trang phù hợp","task"),(0,"Kết thúc","end")]),
("BP02","Đăng ký và điểm danh hoạt động",["Người tham gia","Hệ thống","Cán bộ khoa/Quản trị viên"],[(0,"Xem hoạt động","start"),(1,"Hiển thị hoạt động phù hợp","task"),(0,"Chọn đăng ký","task"),(1,"Kiểm tra điều kiện đăng ký","decision"),(1,"Ghi nhận đăng ký","task"),(0,"Tham gia hoạt động","task"),(2,"Quét QR hoặc nhập MSSV","task"),(1,"Kiểm tra điều kiện điểm danh","decision"),(1,"Ghi nhận đã tham gia và điểm rèn luyện","task"),(0,"Nhận thông báo","end")]),
("BP03","Xét hồ sơ sinh viên hai vòng",["Sinh viên","Hệ thống","Cán bộ lớp","Cán bộ khoa"],[(0,"Tạo hồ sơ","start"),(1,"Tự động kiểm tra tiêu chí","task"),(0,"Bổ sung minh chứng","task"),(0,"Nộp hồ sơ","task"),(1,"Kiểm tra điều kiện nộp","decision"),(2,"Duyệt vòng 1","task"),(1,"Chuyển trạng thái hồ sơ","decision"),(3,"Duyệt vòng 2","task"),(1,"Ghi nhận kết quả cuối","decision"),(0,"Theo dõi kết quả","end")]),
("BP04","Xét hồ sơ cán bộ",["Cán bộ","Hệ thống","Quản trị viên"],[(0,"Tạo hồ sơ","start"),(1,"Tự động kiểm tra tiêu chí và tính điểm","task"),(0,"Bổ sung minh chứng","task"),(0,"Nộp hồ sơ","task"),(1,"Kiểm tra điều kiện nộp","decision"),(2,"Duyệt hồ sơ","task"),(1,"Ghi nhận kết quả","decision"),(0,"Theo dõi kết quả","end")]),
("BP05","Cấu hình bộ tiêu chuẩn và đợt xét",["Quản trị viên","Hệ thống"],[(0,"Tạo hoặc sao chép bộ tiêu chuẩn","start"),(1,"Kiểm tra khả năng chỉnh sửa","decision"),(0,"Cấu hình nhóm, tiêu chí và điều kiện","task"),(1,"Lưu bộ tiêu chuẩn","task"),(0,"Tạo đợt xét","task"),(1,"Kiểm tra loại, thời gian và học kỳ","decision"),(0,"Đồng bộ tiêu chuẩn cho đợt xét","task"),(1,"Tạo bản tiêu chuẩn áp dụng","task"),(0,"Công bố đợt xét","end")]),
("BP06","Nhập kết quả hoạt động và đồng bộ hồ sơ",["Quản trị viên","Hệ thống","Hồ sơ liên quan"],[(0,"Chọn hoạt động và tệp kết quả","start"),(1,"Đọc dữ liệu và đối chiếu người tham gia","decision"),(0,"Xác nhận dữ liệu cần nhập","task"),(1,"Ghi nhận kết quả tham gia và điểm rèn luyện","task"),(2,"Cập nhật lại tiêu chí và điểm hồ sơ","task"),(1,"Ghi nhật ký và cập nhật báo cáo","task"),(0,"Kết thúc","end")]),
]
for code,title,lanes,steps in BUSINESS_EDITABLE:
    business_drawio(code,title,lanes,steps,DRAWIO_BP/f"{code}_{re.sub(r'[^A-Za-z0-9]+','_',title)}.drawio")

# Editable use-case draw.io with <<include>> / <<extend>>
def usecase_drawio(actor, selected, path):
    includes,extends,support=relation_names(selected)
    names=[usecase_name(u) for u in selected]+support
    nodes=[{"id":"actor","label":actor,"lane":0,"kind":"actor","x":30,"y":220,"w":80,"h":110}]
    # One wide lane is enough for editable use-case layout
    lanes=["Mô hình Use Case"]
    for i,name in enumerate(names):
        col=i%3; row=i//3
        nodes.append({"id":f"u{i+1}","label":name,"lane":0,"kind":"start","x":180+col*250,"y":40+row*90,"w":210,"h":55})
    idmap={name:f"u{i+1}" for i,name in enumerate(names)}
    edges=[]
    for u in selected:
        edges.append(("actor",idmap[usecase_name(u)],"","endArrow=none;html=1;"))
    for src,dst in includes:
        edges.append((idmap[src],idmap[dst],"&lt;&lt;include&gt;&gt;","dashed=1;endArrow=open;html=1;"))
    for src,dst in extends:
        edges.append((idmap[src],idmap[dst],"&lt;&lt;extend&gt;&gt;","dashed=1;endArrow=open;html=1;"))
    drawio_doc(f"Use Case - {actor}",lanes,nodes,edges,path)

usecase_drawio("Sinh viên",student_selected,DRAWIO_UC/"02_UseCase_SinhVien.drawio")
usecase_drawio("Cán bộ lớp",class_selected,DRAWIO_UC/"03_UseCase_CanBoLop.drawio")
usecase_drawio("Cán bộ khoa",faculty_selected,DRAWIO_UC/"04_UseCase_CanBoKhoa.drawio")
usecase_drawio("Quản trị viên",admin_selected,DRAWIO_UC/"05_UseCase_Admin.drawio")

guide = """BỘ SƠ ĐỒ CHỈNH SỬA
====================
1) visual-paradigm/*.xmi:
   Import trong Visual Paradigm bằng Project > Import > XMI.
   Các sơ đồ theo vai trò có Association và quan hệ <<include>> / <<extend>>.
2) editable-diagrams/usecase-drawio/*.drawio:
   Mở trực tiếp bằng diagrams.net/draw.io để chỉnh bố cục, nhãn và quan hệ.
3) editable-diagrams/business-process-drawio/*.drawio:
   Sơ đồ nghiệp vụ dạng swimlane có thể chỉnh sửa từng phần tử.
"""
(EDIT/"HUONG_DAN_CHINH_SUA_SO_DO.txt").write_text(guide,encoding="utf-8")

(XMI/"HUONG_DAN_IMPORT_VISUAL_PARADIGM.txt").write_text("""BỘ XMI 2.1 CHO VISUAL PARADIGM
================================
Các file .xmi chứa Actor, Use Case, Association, <<include>> và <<extend>>.
Cách nhập: Project > Import > XMI... Sau khi import, nếu chưa có canvas, tạo Use Case Diagram mới
và kéo các phần tử từ Model Explorer vào sơ đồ rồi Auto Layout.
""",encoding="utf-8")
with zipfile.ZipFile(ROOT/"MeritTrack_VisualParadigm_XMI.zip","w",zipfile.ZIP_DEFLATED) as z:
    for p in sorted(XMI.iterdir()): z.write(p,arcname=p.name)
with zipfile.ZipFile(ROOT/"MeritTrack_Editable_Diagrams.zip","w",zipfile.ZIP_DEFLATED) as z:
    for p in sorted(EDIT.rglob("*")):
        if p.is_file(): z.write(p,arcname=str(p.relative_to(EDIT)))

# ------------------------------------------------------------
# Word helpers
# ------------------------------------------------------------
doc=Document()
sec=doc.sections[0]
sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(2.0); sec.left_margin=Cm(3.0); sec.right_margin=Cm(2.0)
styles=doc.styles
styles["Normal"].font.name="Times New Roman"; styles["Normal"].font.size=Pt(13)
styles["Normal"].paragraph_format.line_spacing=1.3
for name,size in [("Title",18),("Heading 1",16),("Heading 2",14),("Heading 3",13)]:
    st=styles[name]; st.font.name="Times New Roman"; st.font.size=Pt(size); st.font.bold=True

def set_cell(cell,text,bold=False,size=10.5,align=None):
    cell.text=""
    p=cell.paragraphs[0]
    if align is not None: p.alignment=align
    r=p.add_run(str(text)); r.font.name="Times New Roman"; r.font.size=Pt(size); r.bold=bold
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def shade(cell,fill="D9D9D9"):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr()
    tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),"true"); trPr.append(tblHeader)

def add_caption(text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text); r.font.name="Times New Roman"; r.font.size=Pt(11); r.italic=True

def add_picture(path,width=15.5):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path),width=Cm(width))

def add_text(text,bold=False):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    r=p.add_run(text); r.bold=bold; r.font.name="Times New Roman"; r.font.size=Pt(13)
    return p

def add_meta_row(table,label,value):
    c=table.add_row().cells
    m=c[0].merge(c[1])
    p=m.paragraphs[0]
    r=p.add_run(label); r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(10.5)
    r=p.add_run(value); r.font.name="Times New Roman"; r.font.size=Pt(10.5)

def add_flow_rows(table, steps):
    for num,side,action in steps:
        c=table.add_row().cells
        if side=="Actor":
            set_cell(c[0],f"{num}. {action}",False,10)
            set_cell(c[1],"",False,10)
        else:
            set_cell(c[0],"",False,10)
            set_cell(c[1],f"{num}. {action}",False,10)

def add_branch_section(table,title,rows):
    c=table.add_row().cells; m=c[0].merge(c[1])
    set_cell(m,title,True,10.5); shade(m,"D9D9D9")
    if not rows:
        c=table.add_row().cells; m=c[0].merge(c[1]); set_cell(m,"Không có.",False,10)
        return
    add_flow_rows(table, rows)

def add_uc_table(u):
    doc.add_heading(f'{u["id"]}. Use Case {u["name"]}',level=3)
    t=doc.add_table(rows=0,cols=2); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    add_meta_row(t,"- Tên use case: ",f'{u["id"]} – {u["name"]}')
    add_meta_row(t,"- Mô tả sơ lược: ",humanize(u["desc"]))
    add_meta_row(t,"- Actor chính: ",u["actor"])
    add_meta_row(t,"- Actor phụ: ",u["secondary"])
    add_meta_row(t,"- Tiền điều kiện: ",humanize(u["pre"]))
    add_meta_row(t,"- Hậu điều kiện: ",humanize(u["post"]))
    c=t.add_row().cells; m=c[0].merge(c[1]); set_cell(m,"- Luồng sự kiện chính (Main Flow)",True,10.5); shade(m,"D9D9D9")
    c=t.add_row().cells
    set_cell(c[0],"Actor",True,10.5,WD_ALIGN_PARAGRAPH.CENTER); set_cell(c[1],"Hệ thống",True,10.5,WD_ALIGN_PARAGRAPH.CENTER)
    shade(c[0],"E7E6E6"); shade(c[1],"E7E6E6")
    main=flatten_main(u)
    add_flow_rows(t,main)
    add_branch_section(t,"- Luồng sự kiện thay thế (Alternative Flow)",branch_actions(u["alt"],main,"alternative"))
    add_branch_section(t,"- Luồng sự kiện ngoại lệ (Exception Flow)",branch_actions(u["exc"],main,"exception"))
    doc.add_paragraph()

# Footer page number
footer=sec.footer
p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run(); f1=OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"),"begin")
it=OxmlElement("w:instrText"); it.set(qn("xml:space"),"preserve"); it.text="PAGE"
f2=OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"),"end")
r._r.extend([f1,it,f2])

# ------------------------------------------------------------
# Chapter content - exact 8 sections requested
# ------------------------------------------------------------
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("CHƯƠNG 3\nPHÂN TÍCH VÀ THIẾT KẾ HỆ THỐNG")
r.bold=True; r.font.name="Times New Roman"; r.font.size=Pt(18)
add_text("Chương này mô tả các tác nhân, yêu cầu, quy tắc và quy trình nghiệp vụ của MeritTrack; đồng thời trình bày Use Case và đặc tả Use Case bám theo mã nguồn hiện tại.")

# 1
doc.add_heading("3.1. Xác định Actor của hệ thống",level=1)
add_text("MeritTrack có bốn actor nghiệp vụ chính. Role tài khoản và login context được tách biệt, vì vậy cán bộ lớp/cán bộ khoa có thể sử dụng ngữ cảnh người dùng và ngữ cảnh cán bộ trên cùng một tài khoản.")
t=doc.add_table(rows=1,cols=3); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(["STT","Actor","Mô tả"]):
    set_cell(t.rows[0].cells[i],h,True,10.5,WD_ALIGN_PARAGRAPH.CENTER); shade(t.rows[0].cells[i],"E7E6E6")
set_repeat_table_header(t.rows[0])
for i,(a,d) in enumerate(ACTORS,1):
    c=t.add_row().cells; set_cell(c[0],i,False,10,WD_ALIGN_PARAGRAPH.CENTER); set_cell(c[1],a,True,10); set_cell(c[2],d,False,10)

# 2
doc.add_heading("3.2. Sơ đồ Use Case",level=1)
add_text("Sơ đồ tổng quát được tách thêm theo từng actor để tránh mật độ liên kết quá lớn. Bộ XMI 2.1 đính kèm có thể nhập vào Visual Paradigm và tiếp tục chỉnh sửa.")
for img,cap in [(OVERALL,"Hình 3.1. Sơ đồ Use Case tổng quát"),
                (UC_STUDENT,"Hình 3.2. Sơ đồ Use Case – Sinh viên"),
                (UC_CLASS,"Hình 3.3. Sơ đồ Use Case – Cán bộ lớp"),
                (UC_FACULTY,"Hình 3.4. Sơ đồ Use Case – Cán bộ khoa"),
                (UC_ADMIN,"Hình 3.5. Sơ đồ Use Case – Quản trị viên")]:
    add_picture(img,15.5); add_caption(cap)

# 3
doc.add_heading("3.3. Yêu cầu chức năng",level=1)
add_text("Yêu cầu chức năng được liệt kê theo role nhằm làm rõ phạm vi sử dụng và trách nhiệm nghiệp vụ của từng actor.")
for role in ["Sinh viên","Cán bộ lớp","Cán bộ khoa","Quản trị viên"]:
    doc.add_heading(f"3.3.{['Sinh viên','Cán bộ lớp','Cán bộ khoa','Quản trị viên'].index(role)+1}. {role}",level=2)
    rows=[u for u in USECASES if role in uc_role(u)]
    # Avoid listing admin common profile/QR use cases not intended for admin.
    if role=="Quản trị viên":
        rows=[u for u in rows if u["id"] in ["UC01","UC02","UC04","UC07","UC08"] or int(u["id"][2:])>=36]
    if role=="Sinh viên":
        rows=[u for u in rows if int(u["id"][2:])<=20]
    if role=="Cán bộ lớp":
        rows=[u for u in rows if int(u["id"][2:])<=8 or 21<=int(u["id"][2:])<=23 or 29<=int(u["id"][2:])<=35]
    if role=="Cán bộ khoa":
        rows=[u for u in rows if int(u["id"][2:])<=8 or 24<=int(u["id"][2:])<=35]
    tb=doc.add_table(rows=1,cols=4); tb.style="Table Grid"; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(["STT","Mã","Chức năng","Mô tả"]):
        set_cell(tb.rows[0].cells[i],h,True,10,WD_ALIGN_PARAGRAPH.CENTER); shade(tb.rows[0].cells[i],"E7E6E6")
    set_repeat_table_header(tb.rows[0])
    for j,u in enumerate(rows,1):
        c=tb.add_row().cells
        set_cell(c[0],j,False,9.5,WD_ALIGN_PARAGRAPH.CENTER); set_cell(c[1],u["id"],False,9.5,WD_ALIGN_PARAGRAPH.CENTER)
        set_cell(c[2],u["name"],False,9.5); set_cell(c[3],u["desc"],False,9.5)

# 4
doc.add_heading("3.4. Yêu cầu phi chức năng",level=1)
nfr=[
("NFR01","Bảo mật và phân quyền","Kiểm tra quyền ở server; không tin cậy role/context từ client."),
("NFR02","Toàn vẹn dữ liệu","Dùng transaction cho các thao tác cập nhật nhiều bảng quan trọng; duy trì trạng thái nhất quán."),
("NFR03","Bảo toàn lịch sử","Dữ liệu học kỳ cũ, event snapshot, review và timeline phải phục vụ truy vết lịch sử."),
("NFR04","Khả năng truy vết","Ghi audit log cho các thay đổi quản trị nhạy cảm và attendance log cho điểm danh."),
("NFR05","Hiệu năng","Trang danh sách/báo cáo cần lọc theo học kỳ, trạng thái và từ khóa; truy vấn chỉ lấy dữ liệu thuộc phạm vi cần thiết."),
("NFR06","Khả dụng","Giao diện theo role/context nhất quán; thông báo lỗi/thành công rõ ràng; hỗ trợ xem dữ liệu lịch sử."),
("NFR07","An toàn tệp","Minh chứng phải kiểm tra MIME, đuôi tệp, dung lượng và quyền sở hữu submission."),
("NFR08","Khả năng bảo trì","Tách App Router, component, server action, domain service và persistence để giảm coupling."),
]
tb=doc.add_table(rows=1,cols=3); tb.style="Table Grid"; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(["Mã","Nhóm yêu cầu","Nội dung"]):
    set_cell(tb.rows[0].cells[i],h,True,10.5,WD_ALIGN_PARAGRAPH.CENTER); shade(tb.rows[0].cells[i],"E7E6E6")
for row in nfr:
    c=tb.add_row().cells
    for i,v in enumerate(row): set_cell(c[i],v,False,10)

# 5
doc.add_heading("3.5. Quy tắc nghiệp vụ",level=1)
tb=doc.add_table(rows=1,cols=2); tb.style="Table Grid"; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
set_cell(tb.rows[0].cells[0],"Mã",True,10.5,WD_ALIGN_PARAGRAPH.CENTER); set_cell(tb.rows[0].cells[1],"Quy tắc nghiệp vụ",True,10.5,WD_ALIGN_PARAGRAPH.CENTER)
shade(tb.rows[0].cells[0],"E7E6E6"); shade(tb.rows[0].cells[1],"E7E6E6")
for code,rule in BUSINESS_RULES:
    c=tb.add_row().cells; set_cell(c[0],code,True,10); set_cell(c[1],rule,False,10)

# 6
doc.add_heading("3.6. Quy trình nghiệp vụ",level=1)
processes=[
("BP01","Đăng nhập và chuyển ngữ cảnh","Người dùng, MeritTrack","Xác thực tài khoản, tạo session và chuyển context theo role.","Session hợp lệ và dashboard đúng context."),
("BP02","Đăng ký và điểm danh hoạt động","Sinh viên/Cán bộ, Cán bộ khoa/Admin, MeritTrack","Đăng ký activity, tham gia, quét QR/MSSV, xác nhận attended và ghi điểm ĐRL.","Attendance và nguồn điểm được ghi nhận."),
("BP03","Xét hồ sơ sinh viên","Sinh viên, Cán bộ lớp, Cán bộ khoa, MeritTrack","Tạo hồ sơ, auto-evaluate, nộp, duyệt vòng 1, duyệt vòng 2.","Submission kết thúc passed/failed hoặc quay lại trạng thái cần chỉnh sửa."),
("BP04","Xét hồ sơ cán bộ","Cán bộ lớp/Cán bộ khoa, Admin, MeritTrack","Tạo hồ sơ, auto-evaluate, tính điểm, nộp và Admin duyệt một vòng.","Officer submission passed/failed/revise."),
("BP05","Cấu hình bộ tiêu chuẩn và đợt xét","Admin, MeritTrack","Tạo/clone template, cấu hình nhóm/tiêu chí/rules, tạo event và snapshot.","Event có bộ snapshot tiêu chuẩn cố định."),
("BP06","Import kết quả hoạt động và đồng bộ hồ sơ","Admin, MeritTrack","Import attendance, cập nhật ĐRL, audit và re-evaluate hồ sơ liên quan.","Attendance/score/auto results đồng bộ."),
]
tb=doc.add_table(rows=1,cols=5); tb.style="Table Grid"; tb.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,h in enumerate(["Mã","Quy trình","Thành phần tham gia","Tóm tắt","Kết quả"]):
    set_cell(tb.rows[0].cells[i],h,True,9.5,WD_ALIGN_PARAGRAPH.CENTER); shade(tb.rows[0].cells[i],"E7E6E6")
for row in processes:
    c=tb.add_row().cells
    for i,v in enumerate(row): set_cell(c[i],v,False,9)

# 7
doc.add_heading("3.7. Sơ đồ nghiệp vụ",level=1)
for idx,(code,img) in enumerate(BUSINESS_DIAGRAMS,1):
    add_picture(img,15.5)
    title=next(x[1] for x in processes if x[0]==code)
    add_caption(f"Hình 3.{5+idx}. Sơ đồ nghiệp vụ {code} – {title}")

# 8
doc.add_heading("3.8. Đặc tả Use Case",level=1)
add_text("Biểu mẫu đặc tả tuân theo cấu trúc Actor/Hệ thống. Cột Actor chỉ chứa thao tác do người dùng thực hiện; cột Hệ thống chỉ chứa phản ứng/xử lý của MeritTrack. Chỉ số bước chạy liên tục trong Main Flow. Alternative/Exception Flow dùng chỉ số dạng n.x để chỉ rõ nhánh phát sinh từ bước n của luồng chính và ghi rõ điểm quay lại hoặc kết thúc.")
for u in USECASES:
    add_uc_table(u)


# End of Chapter 3 after section 3.8 as requested.
out=ROOT/"Chuong_3_Phan_tich_Thiet_ke_MeritTrack_HoanChinh.docx"
doc.save(out)
print(out)
